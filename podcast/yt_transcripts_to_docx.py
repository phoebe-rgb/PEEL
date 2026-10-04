#!/usr/bin/env python3
"""Tải phụ đề (transcript) YouTube về máy bạn và xuất thành file Word để in, nghe-đọc theo.

Chỉ dùng cho mục đích học cá nhân.

Cài đặt (1 lần):   pip install youtube-transcript-api python-docx
Chạy:              python yt_transcripts_to_docx.py
Mặc định: tự chọn K đoạn "hay nhất" của mỗi video (chấm điểm theo tiếng cười/vỗ tay, câu hỏi,
tốc độ nói, từ thân mật; bỏ phần mở đầu, kết và đoạn quảng cáo).
Tùy chọn:
  --top 3 --minutes 3    số đoạn hay nhất và độ dài mỗi đoạn (phút). Mặc định 3 đoạn x 3 phút
  --mode first           thay vì chọn đoạn hay, lấy từ phút --start đến --end
  --start 0 --end 10     dùng cùng --mode first
  --full                 lấy cả video (rất dài với podcast 1 giờ+)
  --ids ID1 ID2 ...      dùng danh sách video riêng thay cho 24 tập có sẵn
  --out ten_file.docx    tên file xuất ra
Lưu ý: phụ đề tự động của YouTube không có tên người nói và dấu câu có thể sai.
Nếu bị báo lỗi "IP blocked", hãy chạy trên mạng nhà bạn (không dùng VPN/server).
"""
import argparse, re, sys
from youtube_transcript_api import YouTubeTranscriptApi
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK

EPISODES = [
 ("wCoI6G6KLTk","Jay Shetty Interviews His Wife For Valentine's Day"),
 ("J683NQX90r0","Chelsea Handler: ON Spirituality & Meditation"),
 ("Gz3yZC8ZJsA","David Goggins: ON His Relationship With Pain"),
 ("k415svmHwRY","Yuval Harari: ON How To Set Expectations"),
 ("Z2ENRqserkg","Lilly Singh: ON Being Grateful For Depression"),
 ("gjpB9MXCfOE","Mike Posner: ON How Fame Ruined His Life"),
 ("tHMccnWE0Ok","Call Her Daddy – Sonja & Luann"),
 ("y5Wq-BlkdQ0","Call Her Daddy – Harry Jowsey: My Wedding & Divorce"),
 ("ChWjmdZQYZs","Call Her Daddy – Katie Couric"),
 ("pjHzLMfsf3w","Call Her Daddy – Nara Smith"),
 ("XZlbo15aTqE","Call Her Daddy – Alanis Morissette"),
 ("pc3qJ59r3AE","Call Her Daddy – Riley Keough"),
 ("TAetJY3I_Lw","Drew Barrymore Show – Chelsea Handler on dating"),
 ("aXzNNCmrlMI","Drew Barrymore Show – Taylor Tomlinson"),
 ("tphV-tdyW48","Drew Barrymore Show – Dr. Shefali on parenting"),
 ("I2ZfLtDqejY","Drew Barrymore Show – Michael Strahan"),
 ("bAhpBLW_Wr8","Drew Barrymore Show – Lili Reinhart"),
 ("WH4jR_H4UQU","Drew Barrymore Show – Sarah Paulson"),
 ("YmewnMHBOw0","Jimmy Kimmel Live – Amy Schumer"),
 ("yrlod2vMk_A","Jimmy Kimmel Live – Robby Hoffman"),
 ("4nb76mJ2XiU","Jimmy Kimmel Live – Howard Stern"),
 ("nHoquA-kUhI","Jimmy Kimmel Live – Ben Affleck"),
 ("5m4aaVWLGiI","Jimmy Kimmel Live – Billy Crystal"),
 ("s7DcP3r_QHk","Jimmy Kimmel Live – Jamie Foxx"),
]

def fetch(vid, langs):
    try:                                   # youtube-transcript-api >= 1.0
        t = YouTubeTranscriptApi().fetch(vid, languages=langs)
        return [(s.start, s.text) for s in t]
    except AttributeError:                 # phiên bản cũ
        t = YouTubeTranscriptApi.get_transcript(vid, languages=langs)
        return [(s["start"], s["text"]) for s in t]

def paragraphs(snips, start_s, end_s, gap=2.0, maxlen=450):
    paras, cur, cur_t, last_t = [], [], None, None
    for t, text in snips:
        if t < start_s or (end_s is not None and t > end_s):
            continue
        text = re.sub(r"\[.*?\]", "", text).replace("\n", " ").strip()
        if not text:
            continue
        if cur and (t - last_t > gap or sum(len(x) + 1 for x in cur) > maxlen):
            paras.append((cur_t, " ".join(cur))); cur, cur_t = [], None
        if cur_t is None:
            cur_t = t
        cur.append(text); last_t = t
    if cur:
        paras.append((cur_t, " ".join(cur)))
    return paras

INFORMAL = re.compile(r"\b(oh|wow|okay|honestly|literally|totally|actually|you know|i mean|kind of|sort of|"
                      r"like|yeah|crazy|funny|love|hate|obsessed|amazing|insane|wild|weird|hilarious|dude|girl)\b", re.I)
SPONSOR = re.compile(r"sponsor|promo code|discount|subscribe|brought to you|link in the description|use code|"
                     r"patreon|download the app|free trial", re.I)

def best_windows(raw, top=3, win=180, step=20):
    """raw = [(start, text)]. Trả về [(t0, t1)] các cửa sổ điểm cao nhất, không chồng nhau."""
    if not raw:
        return []
    end_t = raw[-1][0]
    cands = []
    t0 = 90.0                                   # bỏ ~1,5 phút mở đầu
    while t0 + win <= end_t - 120:             # bỏ ~2 phút cuối
        seg = [x for x in raw if t0 <= x[0] < t0 + win]
        text = " ".join(x[1] for x in seg)
        if seg and not SPONSOR.search(text):
            words = len(text.split())
            laugh = len(re.findall(r"\[(?:laughter|applause|laughs|cheering)", text, re.I))
            ques = text.count("?")
            inf = len(INFORMAL.findall(text))
            rate = words / win                  # nói nhanh, sôi nổi
            score = laugh * 4 + ques * 1.2 + inf * 0.15 + rate * 5
            cands.append((score, t0))
        t0 += step
    cands.sort(reverse=True)
    picked = []
    for sc, t in cands:
        if all(t + win <= p or t >= p + win for p in picked):
            picked.append(t)
        if len(picked) == top:
            break
    return [(t, t + win) for t in sorted(picked)]

def mmss(s):
    return f"{int(s)//60:02d}:{int(s)%60:02d}"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ids", nargs="*")
    ap.add_argument("--mode", choices=["best", "first"], default="best")
    ap.add_argument("--top", type=int, default=3)
    ap.add_argument("--minutes", type=float, default=3)
    ap.add_argument("--start", type=float, default=0)
    ap.add_argument("--end", type=float, default=10)
    ap.add_argument("--full", action="store_true")
    ap.add_argument("--lang", default="en")
    ap.add_argument("--out", default="YouTube_transcripts.docx")
    a = ap.parse_args()
    eps = [(i, i) for i in a.ids] if a.ids else EPISODES
    end_s = None if (a.full or a.mode == "best") else a.end * 60
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
    doc.add_heading("Script podcast YouTube – dùng để học cá nhân", 0)
    ok = 0
    for n, (vid, title) in enumerate(eps, 1):
        print(f"[{n}/{len(eps)}] {title}")
        try:
            snips = fetch(vid, [a.lang, "en-US", "en-GB"])
        except Exception as e:
            print("   ! bỏ qua:", type(e).__name__, str(e).splitlines()[0][:120])
            continue
        if n > 1 or ok:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(f"Bài {n}: {title}", 1)
        if a.full or a.mode == "first":
            spans = [(a.start * 60, end_s)]
            desc = "cả video" if a.full else f"phút {a.start:g}–{a.end:g}"
        else:
            spans = best_windows(snips, a.top, a.minutes * 60) or [(0, a.minutes * 60 * a.top)]
            desc = f"{len(spans)} đoạn hay nhất"
        doc.add_paragraph(f"Link: https://youtu.be/{vid}   |   {desc}")
        for k, (s0, s1) in enumerate(spans, 1):
            if len(spans) > 1:
                doc.add_heading(f"Đoạn {k}: {mmss(s0)} – {mmss(s1)}", 2)
            for t, text in paragraphs(snips, s0, s1):
                p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
                r = p.add_run(f"[{mmss(t)}] "); r.bold = True; r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
                p.add_run(text)
        ok += 1
    doc.save(a.out)
    print(f"Xong: {ok}/{len(eps)} video -> {a.out}")

if __name__ == "__main__":
    sys.exit(main())
